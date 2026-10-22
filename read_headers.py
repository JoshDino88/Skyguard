# -*- coding: utf-8 -*-
import openpyxl, sys, io
sys.stdout = io.TextIOWrapper(sys.stdout.buffer, encoding='utf-8', errors='replace')
p = r"C:\Users\Echo Li\Documents\deepseek-harness\default-workspace\skyguard\data\Track 3 Data.xlsx"
wb = openpyxl.load_workbook(p, data_only=True)
for ws in wb.worksheets:
    hdr = [c.value for c in ws[1]]
    print(ws.title, "| COLS:", len(hdr))
    print("  " + " | ".join("" if h is None else str(h) for h in hdr))
    print("  max_row:", ws.max_row)
